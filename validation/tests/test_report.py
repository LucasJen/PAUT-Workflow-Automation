"""The Word report of a run: every section, every check, pictures, and the page's download."""
import os
import shutil
import tempfile

from docx import Document
from django.test import TestCase
from django.urls import reverse

from analysis.tests import builders
from validation import runner
from validation.kit import Kit
from validation.report import build


class ReportTests(TestCase):
    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.folder, True)
        source = os.path.join(self.folder, 'source')
        os.makedirs(source)
        self.kit = Kit(os.path.join(self.folder, 'kit'))
        self.kit.add(builders.weld_file(source), ['decoding'], 'synthetic weld')
        self.run = runner.execute(runner.create(['L0', 'L2'], self.kit.root))

    def test_the_report_has_every_section_and_check(self):
        doc = Document(build(self.run))
        text = '\n'.join(p.text for p in doc.paragraphs)
        for heading in ('Analysis page validation', 'Summary', 'Scope and limitations', 'Evidence: the kit', 'L0 Decoding',
                        'L2 Sizing and detection', 'Review and approval'):
            self.assertIn(heading, text)
        self.assertIn(f'Run #{self.run.pk}: PASSED', text)
        self.assertIn('L2 Sizing and detection: not built yet - nothing is claimed for it.', text)
        cells = {c.text for t in doc.tables for r in t.rows for c in r.cells}
        self.assertIn('nde/weld.nde', cells)
        self.assertIn(self.run.manifest_sha, cells)

    def test_pictures_go_in(self):
        from PIL import Image
        folder = runner.run_folder(self.run.pk)
        os.makedirs(folder, exist_ok=True)
        self.addCleanup(shutil.rmtree, folder, True)
        Image.new('RGB', (40, 20), 'red').save(os.path.join(folder, 'shot.png'))
        self.run.results['layers'][0]['checks'][0]['pictures'] = [{'name': 'shot.png', 'caption': 'A picture'}]
        self.run.save()
        doc = Document(build(self.run))
        self.assertEqual(len(doc.inline_shapes), 1)

    def test_the_page_downloads_it(self):
        response = self.client.get(reverse('validation-run-report', args=[self.run.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response['Content-Disposition'])
