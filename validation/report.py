"""
A validation run as a Word document for the Level III: what was run and on what (app version, the
kit's files and fingerprints), the verdict by layer, what isn't covered, every check with expected
against measured and the rule, the pictures, and a sign-off page.
"""
import io
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .kit import ROLES
from .runner import run_folder

GREEN, RED, GREY = RGBColor(0x06, 0x76, 0x47), RGBColor(0xB4, 0x23, 0x18), RGBColor(0x66, 0x70, 0x85)
LIMITS = [
    'TFM / FMC groups are not shown by the app (listed under L0): use OmniPC.',
    'Pipe curvature corrections, TCG / DAC curves and encoder-free (time-based) scans are not validated.',
    "Readings that depend on cursor positions or a zone (S/U/I(m-r), TminZ) are compared only where the export gives "
    'the positions; otherwise they are listed as not compared.',
    "Grading (L5) follows the case table written from 100-UT-020 Rev 9; the table's expected grades are for the Level III "
    'to confirm. The app cannot tell a crack, lack of fusion or incomplete penetration from other indications: under B31.1 '
    'such indications are flagged for the technician to characterise.',
    'The checks cover the files in the kit; other probes, wedges, materials and scan types need their own files in the kit.',
]


def _shade(cell, hex_colour):
    tc = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), hex_colour)
    tc.append(shading)


def _table(doc, header, rows, widths=None, size=8):
    """A table; header None: name / value rows, the names shaded."""
    columns = len(header) if header else len(rows[0])
    table = doc.add_table(rows=1 if header else 0, cols=columns)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells if header else [], header or []):
        cell.text = ''
        run = cell.paragraphs[0].add_run(text)
        run.bold, run.font.size = True, Pt(size)
        _shade(cell, 'E8ECF2')
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ''
            colour = None
            if isinstance(value, tuple):
                value, colour = value
            run = cell.paragraphs[0].add_run(str(value))
            run.font.size = Pt(size)
            if colour is not None:
                run.font.color.rgb = colour
                run.bold = True
        if not header:
            _shade(cells[0], 'E8ECF2')
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    return table


def _result(ok):
    return ('-', GREY) if ok is None else (('Pass', GREEN) if ok else ('FAIL', RED))


def _pictures(doc, run, pictures):
    paths = [(os.path.join(run_folder(run.pk), p['name']), p.get('caption', '')) for p in pictures]
    paths = [(p, c) for p, c in paths if os.path.isfile(p)]
    if not paths:
        return
    for path, caption in paths:   # one under the other, full width: OmniPC's and the app's at the same size
        doc.add_picture(path, width=Inches(6.6))
        doc.add_paragraph(caption).runs[0].font.size = Pt(8)


def build(run):
    """The report as a BytesIO (.docx)."""
    results = run.results or {}
    layers = results.get('layers') or []
    doc = Document()
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Calibri', Pt(10)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.8)
        section.top_margin = section.bottom_margin = Inches(0.7)

    doc.add_heading('Analysis page validation', 0)
    verdict = {'passed': 'PASSED', 'failed': 'FAILED', 'error': 'STOPPED BY AN ERROR', 'cancelled': 'STOPPED'}.get(run.status, run.status.upper())
    p = doc.add_paragraph()
    r = p.add_run(f'Run #{run.pk}: {verdict} - {run.passed} checks passed, {run.failed} failed')
    r.bold, r.font.size = True, Pt(14)
    r.font.color.rgb = GREEN if run.status == 'passed' else RED
    doc.add_paragraph(
        'The Analysis page reads phased-array data from Evident .nde files and shows it like OmniPC. This run checks it against '
        'sources that share none of its code: the raw data read by an independent reader (L0), OmniPC\'s own exports (L1, L2), '
        'calibration blocks with known reflectors (L3), the data under each pixel of the real page (L4), the acceptance '
        'criteria as written (L5), and the results approved before (L6). Every evidence file is fingerprinted (SHA-256): '
        'a file that changed after it was added fails and is not used.')
    _table(doc, None, [
        ['Started', run.created_at.strftime('%Y-%m-%d %H:%M')],
        ['Finished', run.finished_at.strftime('%Y-%m-%d %H:%M') if run.finished_at else '-'],
        ['App version (git commit)', run.version or 'unknown'],
        ['Computer', run.computer or '-'],
        ['Kit folder', run.kit_path],
        ['Kit fingerprint (manifest SHA-256)', run.manifest_sha or '-'],
        ['Layers run', ', '.join(run.layers)],
    ], widths=[2.2, 4.7], size=9)

    doc.add_heading('Summary', 1)
    def layer_result(layer):
        if layer['status'] == 'not built':
            return 'Not built', GREY
        if layer['status'] == 'error':
            return 'Error', RED
        return _result(not layer['failed'] if layer['checks'] else None)
    _table(doc, ['Layer', 'What it proves', 'Checked against', 'Passed', 'Failed', 'Result'], [
        [layer['id'] + ' ' + layer['title'], layer['question'], layer['against'], layer['passed'], layer['failed'], layer_result(layer)]
        for layer in layers], widths=[1.2, 2.6, 1.5, 0.5, 0.5, 0.6])
    kit_files = results.get('kit', {}).get('files') or []
    bad = [f for f in kit_files if not f['ok']]
    doc.add_paragraph(f"Kit fingerprints: {len(kit_files) - len(bad)} of {len(kit_files)} files match." +
                      (' Not used: ' + '; '.join(f"{f['path']} ({f['reason']})" for f in bad) if bad else ''))

    doc.add_heading('Scope and limitations', 1)
    for text in LIMITS:
        doc.add_paragraph(text, style='List Bullet')
    for layer in layers:
        if layer['status'] == 'not built':
            doc.add_paragraph(f"{layer['id']} {layer['title']}: not built yet - nothing is claimed for it.", style='List Bullet')
        elif layer['message'] and not layer['checks']:
            doc.add_paragraph(f"{layer['id']} {layer['title']}: {layer['message']}", style='List Bullet')
    not_shown = [c['subject'] for layer in layers for c in layer['checks'] if c['kind'] == 'info' and layer['id'] == 'L0']
    if not_shown:
        doc.add_paragraph('Groups the app does not show: ' + '; '.join(not_shown), style='List Bullet')

    doc.add_heading('Evidence: the kit', 1)
    from .kit import Kit
    try:
        manifest = Kit(run.kit_path).files
    except Exception:   # the kit may have moved since: the run still reports what it checked
        manifest = []
    _table(doc, ['File', 'Used for', 'Source', 'SHA-256'], [
        [e['path'], ', '.join(ROLES.get(r, r) for r in e.get('roles') or []), e.get('source', ''), e.get('sha256', '')]
        for e in manifest], widths=[1.8, 1.3, 1.7, 2.1], size=7)

    for layer in layers:
        if layer['checks']:   # a layer with nothing to show stays on the page before
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(f"{layer['id']} {layer['title']}", 1)
        doc.add_paragraph(f"{layer['question']} Checked against: {layer['against']}.")
        if layer.get('message'):
            doc.add_paragraph(layer['message']).runs[0].font.color.rgb = GREY
        checks = sorted(layer['checks'], key=lambda c: c['passed'])   # failures first
        compact = checks and all(len(c['items']) <= 2 and not c.get('pictures') for c in checks)
        if compact:
            _table(doc, ['Subject', 'Compared', 'Expected', 'Measured', 'Rule', 'Result'], [
                [c['subject'], i['what'], i['expected'], i['measured'], i['rule'], _result(i['ok'])]
                for c in checks for i in c['items']], widths=[1.6, 1.0, 1.0, 1.0, 1.8, 0.5], size=7)
            continue
        for c in checks:
            heading = doc.add_heading(f"{c['subject']} - {c['title']}", 3)
            if c['kind'] == 'check' and not c['passed']:
                for r in heading.runs:
                    r.font.color.rgb = RED
            _table(doc, ['Compared', f"Expected ({c['against']})", 'Measured (the app)', 'Rule', 'Result'],
                   [[i['what'], i['expected'], i['measured'], i['rule'], _result(i['ok'])] for i in c['items']],
                   widths=[1.7, 1.5, 1.5, 1.6, 0.5], size=7)
            if c.get('note'):
                doc.add_paragraph(c['note']).runs[0].font.size = Pt(8)
            _pictures(doc, run, c.get('pictures') or [])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading('Review and approval', 1)
    doc.add_paragraph('I have reviewed this validation run, its evidence and its limitations.')
    _table(doc, None, [
        ['Reviewed by (Level III)', ''], ['Certification / number', ''], ['Approved for use', 'Yes  /  No'],
        ['Conditions or limitations', '\n\n'], ['Signature', '\n'], ['Date', ''],
    ], widths=[2.2, 4.7], size=10)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out
