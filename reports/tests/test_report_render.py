"""Renders the real Word template with a populated report and inspects the output."""
import datetime
import io
import re
import shutil
import tempfile
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from docx import Document
from PIL import Image

from reports.models import Report, ReportImage, ResultsRow, ResultsTable, Setup
from reports.services.report_render import render_report, with_unit

TAG = re.compile(r'\{\{|\}\}|\{%|%\}')


def png_bytes(color='blue'):
    buf = io.BytesIO()
    Image.new('RGB', (40, 20), color).save(buf, 'PNG')
    return buf.getvalue()


class RenderTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.media = tempfile.mkdtemp()
        cls.override = override_settings(MEDIA_ROOT=cls.media)
        cls.override.enable()

    @classmethod
    def tearDownClass(cls):
        cls.override.disable()
        shutil.rmtree(cls.media, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.report = Report.objects.create(
            document_filename='15V3', document_title='PAUT examination of Deethanizer Absorber 15V3',
            client='ACME Refining', location='Rosemount, Minnesota', work_order='WO123',
            report_date=datetime.date(2026, 9, 3), test_date=datetime.date(2026, 8, 13),
            project_type='Level A HIC', procedure='100-UT-031 Rev. 1\n100-UT-021 Rev 9.0',
            technician_name='Pat Tech', certification='Ultrasonic Level II',
            assistant_name='Sam Rope', assistant_certification='Rope Access II',
            examination_scope='Scope paragraph one.\n\nScope paragraph two.',
            executive_summary='Results summary.', equipment_id='15V3',
            equipment_overview='Access by rope.', work_scope='Encoded HydroFORM scans.',
            x_axis_reference='from north', y_axis_reference='from CW1',
            ut_method='HydroFORM 0-degree PAUT\nPAUT Angle Beam',
        )
        Setup.objects.create(report=self.report, order=0, beam_formation='HydroFORM', scope_platform='OmniScan',
                             scope_model='X3 64', x_res='0.039', foc_depth='0.700', freq='7.5',
                             wave_propagation='Longitudinal', material_temp='80', tr_min='0.100', tr_max='1.500"')
        Setup.objects.create(report=self.report, order=1, beam_formation='Angle Beam', transducer_model='10L32-A1')
        table = ResultsTable.objects.create(report=self.report, columns=['Scan ID', 'Orient', 'X', 'Y', 'Avg', 'Min', 'Results'])
        for i, (scan, min_thk) in enumerate([('CW1 Top', '0.800'), ('CW4 Bottom', '0.693'), ('LS2', 'N/A')]):
            ResultsRow.objects.create(table=table, order=i, cells=[scan, 'Circ', '14 - 103', '0 - 2.4', '0.860', min_thk, f'{scan} comments'])
        ReportImage.objects.create(report=self.report, kind=ReportImage.DRAWING, caption='File Drawing',
                                   image=SimpleUploadedFile('drawing.png', png_bytes()))
        # Photo summary: one image tied to a results row, one standalone snip
        ReportImage.objects.create(report=self.report, kind=ReportImage.SCAN, scan_id='CW4 Bottom', order=0,
                                   image=SimpleUploadedFile('cw4.png', png_bytes('red')))
        ReportImage.objects.create(report=self.report, kind=ReportImage.SCAN, caption='Extra snip', order=1,
                                   image=SimpleUploadedFile('extra.png', png_bytes('green')))

        self.content = render_report(self.report)
        self.doc = Document(io.BytesIO(self.content))
        self.zip = zipfile.ZipFile(io.BytesIO(self.content))

    def all_text(self):
        parts = [n for n in self.zip.namelist() if re.match(r'word/(document|header\d+|footer\d+)\.xml', n)]
        return '\n'.join(re.sub(r'<[^>]+>', '', self.zip.read(n).decode()) for n in parts)

    def test_no_template_tags_left_anywhere(self):
        self.assertIsNone(TAG.search(self.all_text()))

    def test_cover_and_headers(self):
        text = self.all_text()
        self.assertIn('PAUT EXAMINATION OF DEETHANIZER ABSORBER 15V3', text)
        self.assertIn('3 September, 2026', text)  # cover date and footer
        self.assertIn('8/13/2026', text)
        headers = ''.join(re.sub(r'<[^>]+>', '', self.zip.read(n).decode()) for n in self.zip.namelist() if 'header' in n)
        self.assertIn('ACME Refining', headers)
        self.assertNotIn('Flint Hills', text)

    def test_personnel_and_procedures(self):
        text = self.all_text()
        for expected in ('Pat Tech', 'Sam Rope', 'Rope Access II', '100-UT-031 Rev. 1', '100-UT-021 Rev 9.0'):
            self.assertIn(expected, text)

    def test_one_equipment_section_per_setup(self):
        headings = [p.text for p in self.doc.paragraphs if p.text.startswith('Equipment Details:')]
        self.assertEqual(headings, ['Equipment Details: HydroFORM', 'Equipment Details: Angle Beam'])

    def test_units_added_once(self):
        text = self.all_text()
        self.assertIn('0.700"', text)
        self.assertIn('80°F', text)
        self.assertIn('1.500"', text)
        self.assertNotIn('1.500""', text)  # value already had the inch mark

    def test_results_table_rows_and_min_highlight(self):
        table = next(t for t in self.doc.tables if t.rows and t.rows[0].cells[0].text.startswith('PAUT'))
        self.assertEqual(table.rows[0].cells[0].text.split('\n')[0], 'PAUT HydroFORM Work Scope')
        data = [r.cells[0].text for r in table.rows[2:]]
        self.assertEqual(data, ['CW1 Top', 'CW4 Bottom', 'LS2'])
        highlighted = [run.text for row in table.rows[2:] for p in row.cells[5].paragraphs for run in p.runs
                       if run.font.highlight_color is not None]
        self.assertEqual(highlighted, ['0.693'])

    def paragraph_styles(self):
        return [(p.style.name, p.text) for p in self.doc.paragraphs]

    def test_only_equipment_drawings_under_drawing_heading(self):
        headings = [text for style, text in self.paragraph_styles() if style == 'Heading 2' and not text.startswith('Equipment')]
        self.assertEqual(headings, ['File Drawing'])

    def test_photo_summary_blocks_use_results_comments(self):
        blocks = [t for t in self.doc.tables if len(t.rows) == 2 and t.rows[1].cells[-1].text.startswith('Comments')]
        summary = [(t.rows[1].cells[0].text.strip(), t.rows[1].cells[-1].paragraphs[-1].text) for t in blocks]
        self.assertEqual(summary, [('CW4 Bottom', 'CW4 Bottom comments'), ('Extra snip', '')])
        for t in blocks:
            self.assertTrue(t.rows[0].cells[0]._tc.xpath('.//pic:pic'), 'scan image missing')

    def test_techniques_and_images_embedded(self):
        texts = [p.text for p in self.doc.paragraphs]
        self.assertIn('HydroFORM 0-degree PAUT', texts)
        self.assertIn('PAUT Angle Beam', texts)
        media = [n for n in self.zip.namelist() if n.startswith('word/media/')]
        self.assertEqual(len(media), 6)  # 3 template images + 1 drawing + 2 scan images

    def test_multi_paragraph_text(self):
        texts = [p.text for t in self.doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
        self.assertIn('Scope paragraph one.', texts)
        self.assertIn('Scope paragraph two.', texts)

    def test_word_refreshes_fields_on_open(self):
        settings_xml = self.zip.read('word/settings.xml').decode()
        self.assertIn('w:updateFields w:val="true"', settings_xml)

    def test_empty_optional_sections_are_dropped(self):
        texts = [p.text for p in self.doc.paragraphs]
        self.assertNotIn('DATA COMPARISON', texts)  # no comparison figures yet


class WithUnitTests(TestCase):
    def test_with_unit(self):
        self.assertEqual(with_unit('0.500', '"'), '0.500"')
        self.assertEqual(with_unit('0.500"', '"'), '0.500"')
        self.assertEqual(with_unit('N/A', '"'), 'N/A')
        self.assertEqual(with_unit('', '"'), '')
