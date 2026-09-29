"""In-app preview (docx-preview page + inline .docx) and master-template section switches."""
import io
import os
import re
import tempfile
import zipfile
from unittest import mock

from django.test import TestCase, override_settings
from django.urls import reverse
from docx import Document

from reports.models import Report, ReportImage, ReportPerson, ResultsRow, ResultsTable, Setup
from reports.report_types import HIC_RESULTS_COLUMNS, SECTIONS, ReportType
from reports.services.report_render import render_report
from reports.tests.test_create_report import post_data
from reports.tests.test_phase2 import MediaMixin, png

TAG = re.compile(r'\{\{|\}\}|\{%|%\}')


def report_with_setup(**fields):
    report = Report.objects.create(document_filename='15V3', **fields)
    Setup.objects.create(report=report, title='HydroFORM')
    return report


class PreviewViewTests(TestCase):
    def test_preview_page_loads_renderer_and_points_at_inline_docx(self):
        report = report_with_setup()
        resp = self.client.get(reverse('preview-report', args=[report.pk]))
        self.assertContains(resp, f'data-docx-url="{reverse("report-docx", args=[report.pk])}"')
        self.assertContains(resp, 'vendor/docx-preview/docx-preview.min.js')
        self.assertContains(resp, 'vendor/docx-preview/jszip.min.js')

    def test_preview_without_setups_returns_to_editor(self):
        report = Report.objects.create()
        resp = self.client.get(reverse('preview-report', args=[report.pk]), follow=True)
        self.assertContains(resp, 'Add at least one UT setup')

    def test_inline_docx_is_valid_and_writes_no_server_copy(self):
        report = report_with_setup()
        with tempfile.TemporaryDirectory() as out, override_settings(REPORT_OUTPUT_DIR=out):
            resp = self.client.get(reverse('report-docx', args=[report.pk]))
            self.assertEqual(os.listdir(out), [])
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn('attachment', resp.get('Content-Disposition', ''))
        self.assertEqual(resp['Cache-Control'], 'no-store')
        Document(io.BytesIO(b''.join(resp.streaming_content)))  # opens as a Word document

    def test_save_and_preview(self):
        report = report_with_setup()
        data = post_data(report=report, setups=[{'id': report.setups.get().pk, 'report': report.pk, 'title': 'HydroFORM'}])
        data['preview'] = ''
        resp = self.client.post(reverse('create-report'), data)
        self.assertRedirects(resp, reverse('preview-report', args=[report.pk]))

    def test_report_list_links_to_preview(self):
        report = report_with_setup()
        self.assertContains(self.client.get(reverse('report-list')), reverse('preview-report', args=[report.pk]))


class MasterTemplateTests(MediaMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.report = report_with_setup(equipment_id='15V3', x_axis_reference='from north')
        table = ResultsTable.objects.create(report=self.report, columns=[h for _, h in HIC_RESULTS_COLUMNS])
        ResultsRow.objects.create(table=table, order=0, cells=['CW1 Top', 'Circ', '1', '2', '0.8', '0.7', 'ok'])
        ReportImage.objects.create(report=self.report, kind=ReportImage.DRAWING, caption='FILE DRAWING', image=png('d.png'))
        ReportImage.objects.create(report=self.report, kind=ReportImage.SCAN, scan_id='CW1 Top', image=png('s.png', 'blue'))
        ReportPerson.objects.create(report=self.report, name='Pat', prepared=True)

    def render_as(self, sections):
        report_type = ReportType('test', 'Test', sections=sections, results_columns=HIC_RESULTS_COLUMNS)
        with mock.patch('reports.services.report_render.get_report_type', return_value=report_type):
            content = render_report(self.report)
        text = re.sub(r'<[^>]+>', '', re.sub(r'</w:p>', '\n', zipfile.ZipFile(io.BytesIO(content)).read('word/document.xml').decode()))
        return Document(io.BytesIO(content)), text

    def headings(self, doc):
        return [p.text for p in doc.paragraphs if p.style.name == 'Heading 1']

    def test_all_sections(self):
        doc, text = self.render_as(SECTIONS)
        self.assertEqual(self.headings(doc), ['INTRODUCTION', 'DISCUSSION', 'DRAWING', 'Calibrations', 'Results', 'SCAN IMAGES'])
        self.assertIsNone(TAG.search(text))

    def test_type_without_results_images_or_drawings(self):
        doc, text = self.render_as(tuple(s for s in SECTIONS if s not in ('results', 'images', 'drawings')))
        self.assertEqual(self.headings(doc), ['INTRODUCTION', 'DISCUSSION', 'Calibrations'])
        self.assertIsNone(TAG.search(text))
        self.assertEqual(len(doc.sections), 4)  # no leftover empty section after Calibrations

    def test_type_with_results_but_no_scan_images_has_no_blank_last_section(self):
        doc, _ = self.render_as(tuple(s for s in SECTIONS if s != 'images'))
        self.assertEqual(self.headings(doc)[-1], 'Results')
        self.assertEqual(len(doc.sections), 4)

    def test_toc_placeholder_instead_of_stale_entries(self):
        doc, text = self.render_as(SECTIONS)
        self.assertIn('Table of contents – updated by Word when the document is opened', text)
        self.assertNotIn('NOZZLE LAYOUT', text)  # entry from the reference report's cached TOC
