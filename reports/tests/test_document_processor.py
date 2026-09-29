import os
import tempfile
from types import SimpleNamespace

from django.test import SimpleTestCase
from docx import Document
from PIL import Image

from reports.services.document_processor import WordTemplateProcessor


class InsertImagesTests(SimpleTestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = tmp.name

        doc = Document()
        doc.add_paragraph('Before')
        doc.add_paragraph('{{IMAGE_BLOCK}}')
        doc.add_paragraph('After')
        self.template = os.path.join(self.dir, 'template.docx')
        doc.save(self.template)

    def make_image(self, name, caption=''):
        path = os.path.join(self.dir, name)
        Image.new('RGB', (10, 10), 'red').save(path)
        return SimpleNamespace(image=SimpleNamespace(path=path, name=name), caption=caption)

    def process(self, images):
        output = os.path.join(self.dir, 'out', 'result.docx')
        processor = WordTemplateProcessor(self.template, output)
        processor.insert_images(images)
        processor.save()
        return Document(output)

    def test_all_images_are_inserted_in_order_with_captions(self):
        doc = self.process([
            self.make_image('a.png', 'Caption A'),
            self.make_image('b.png'),
            self.make_image('c.png', 'Caption C'),
        ])

        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(texts, ['Before', '', 'Caption A', '', '', 'Caption C', 'After'])
        self.assertEqual(len(doc.inline_shapes), 3)
        self.assertNotIn('{{IMAGE_BLOCK}}', '\n'.join(texts))

    def test_no_images_clears_sentinel(self):
        doc = self.process([])
        self.assertNotIn('{{IMAGE_BLOCK}}', '\n'.join(p.text for p in doc.paragraphs))
