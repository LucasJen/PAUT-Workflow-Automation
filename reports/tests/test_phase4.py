"""Prose fields, text library, technique bullets and Duplicate for repeat inspections."""
import datetime
import io
import re
import zipfile

from django.test import TestCase
from django.urls import reverse
from docx import Document

from reports.models import (
    Report, ReportImage, ReportPerson, ResultsRow, ResultsTable, Setup, SetupImage, TextSnippet,
)
from reports.services.report_render import render_report
from reports.tests.test_phase2 import MediaMixin, png


def bullets(doc):
    return [p for p in doc.paragraphs if p.style.name == 'Bullets']


class TextLibraryTests(TestCase):
    def test_seeded_from_reference_report(self):
        names = set(TextSnippet.objects.values_list('kind', 'name'))
        self.assertLessEqual({('technique', 'HydroFORM'), ('technique', 'Angle Beam'), ('technique', 'TFM'),
                              ('discussion', 'default')}, names)

    def test_list_new_edit_delete(self):
        self.assertContains(self.client.get(reverse('snippet-list')), 'ENCODED HydroFORM 0-degree PAUT')
        resp = self.client.post(reverse('new-snippet'), {'kind': 'technique', 'name': 'Corrosion Mapping',
                                                         'title': 'Encoded Corrosion Mapping', 'body': 'Maps thickness.'})
        self.assertRedirects(resp, reverse('snippet-list'))
        snippet = TextSnippet.objects.get(name='Corrosion Mapping')
        self.client.post(reverse('edit-snippet', args=[snippet.pk]),
                         {'kind': 'technique', 'name': 'Corrosion Mapping', 'title': 'CM', 'body': 'Updated.'})
        snippet.refresh_from_db()
        self.assertEqual(snippet.body, 'Updated.')
        self.client.post(reverse('edit-snippet', args=[snippet.pk]), {'delete': ''})
        self.assertFalse(TextSnippet.objects.filter(pk=snippet.pk).exists())

    def test_name_unique_per_kind(self):
        resp = self.client.post(reverse('new-snippet'), {'kind': 'technique', 'name': 'HydroFORM', 'title': 'x', 'body': 'y'})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(TextSnippet.objects.filter(kind='technique', name='HydroFORM').count(), 1)

    def test_nav_link(self):
        self.assertContains(self.client.get(reverse('home')), f'href="{reverse("snippet-list")}"')


class ReportTextTests(TestCase):
    def render(self, report):
        return Document(io.BytesIO(render_report(report)))

    def test_technique_bullets_follow_setup_titles(self):
        report = Report.objects.create(ut_method='ignored when setups have titles')
        for i, title in enumerate(['hydroform', 'TFM', 'HydroFORM', 'Weld Scan']):  # case-insensitive, deduplicated
            Setup.objects.create(report=report, order=i, title=title)
        paras = bullets(self.render(report))
        self.assertEqual(len(paras), 3)
        hydro = paras[0]
        self.assertTrue(hydro.text.startswith('ENCODED HydroFORM 0-degree PAUT – HydroFORM utilizes'))
        bold = ''.join(r.text for r in hydro.runs if r.bold)
        self.assertEqual(bold, 'ENCODED HydroFORM 0-degree PAUT')
        self.assertTrue(paras[1].text.startswith('Total Focusing Method (TFM) – '))
        self.assertEqual(paras[2].text, 'Weld Scan')  # no library entry: just the title
        self.assertNotIn('ignored', '\n'.join(p.text for p in paras))

    def test_ut_method_fallback_without_titles(self):
        report = Report.objects.create(ut_method='Manual UT\nThickness readings')
        Setup.objects.create(report=report)
        self.assertEqual([p.text for p in bullets(self.render(report))], ['Manual UT', 'Thickness readings'])

    def test_discussion_standard_or_own(self):
        report = Report.objects.create()
        Setup.objects.create(report=report)
        texts = [p.text for p in self.render(report).paragraphs]
        self.assertTrue(any(t.startswith('In the Results section, a table has been constructed') for t in texts))

        report.discussion = 'Our own discussion.\n\nSecond paragraph.'
        report.save()
        texts = [p.text for p in self.render(report).paragraphs]
        self.assertIn('Our own discussion.', texts)
        self.assertIn('Second paragraph.', texts)
        self.assertFalse(any(t.startswith('In the Results section, a table has been') for t in texts))

    def test_asset_description_opens_the_introduction(self):
        report = Report.objects.create(asset_description='15V-3 is a vertical vessel in the 15 Unit.')
        Setup.objects.create(report=report)
        texts = [p.text for p in self.render(report).paragraphs]
        intro = texts.index('INTRODUCTION')
        self.assertEqual(texts[intro + 1], '15V-3 is a vertical vessel in the 15 Unit.')

    def test_editor_shows_new_fields(self):
        resp = self.client.get(reverse('create-report'))
        self.assertContains(resp, 'name="asset_description"')
        self.assertContains(resp, 'name="discussion"')
        self.assertContains(resp, 'Access &amp; surface condition')
        self.assertContains(resp, 'data-section="discussion"')


class DuplicateTests(MediaMixin, TestCase):
    def test_duplicate_carries_setup_people_drawings_and_text_but_not_results(self):
        original = Report.objects.create(
            document_filename='15V3 2026', client='FHR', asset_description='Vessel.', discussion='Disc.',
            report_date=datetime.date(2026, 9, 3), test_date=datetime.date(2026, 8, 13), test_end_date=datetime.date(2026, 8, 25))
        ReportPerson.objects.create(report=original, name='Pat', prepared=True)
        setup = Setup.objects.create(report=original, title='HydroFORM', procedure='P-1')
        SetupImage.objects.create(setup=setup, image=png('cal.png'))
        ReportImage.objects.create(report=original, kind=ReportImage.DRAWING, caption='FILE DRAWING', image=png('d.png', 'blue'))
        ReportImage.objects.create(report=original, kind=ReportImage.SCAN, scan_id='CW1', image=png('s.png', 'green'))
        table = ResultsTable.objects.create(report=original, columns=['Scan ID'])
        ResultsRow.objects.create(table=table, cells=['CW1'])

        resp = self.client.post(reverse('report-list'), {'selected': [original.pk], 'duplicate': ''})
        dup = Report.objects.exclude(pk=original.pk).get()
        self.assertRedirects(resp, f"{reverse('create-report')}?loaded={dup.pk}", fetch_redirect_response=False)

        self.assertEqual(dup.document_filename, '15V3 2026 (copy)')
        self.assertEqual((dup.client, dup.asset_description, dup.discussion), ('FHR', 'Vessel.', 'Disc.'))
        self.assertEqual((dup.report_date, dup.test_date, dup.test_end_date), (None, None, None))
        self.assertEqual(list(dup.people.values_list('name', 'prepared')), [('Pat', True)])
        dup_setup = dup.setups.get()
        self.assertEqual((dup_setup.title, dup_setup.procedure), ('HydroFORM', 'P-1'))
        self.assertEqual(dup_setup.images.count(), 1)
        self.assertEqual(list(dup.images.values_list('kind', 'caption')), [('drawing', 'FILE DRAWING')])
        self.assertFalse(hasattr(dup, 'results_table') and ResultsTable.objects.filter(report=dup).exists())
        # The original is untouched
        self.assertEqual(original.setups.count(), 1)
        self.assertEqual(original.images.count(), 2)
        self.assertTrue(ResultsTable.objects.filter(report=original).exists())

    def test_duplicate_renders_cleanly(self):
        original = Report.objects.create(document_filename='r')
        Setup.objects.create(report=original, title='TFM')
        self.client.post(reverse('report-list'), {'selected': [original.pk], 'duplicate': ''})
        dup = Report.objects.exclude(pk=original.pk).get()
        xml = zipfile.ZipFile(io.BytesIO(render_report(dup))).read('word/document.xml').decode()
        self.assertIsNone(re.search(r'\{\{|\{%', xml))
