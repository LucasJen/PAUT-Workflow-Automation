"""Personnel list, test date range and data-comparison figures."""
import datetime
import io

from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from docx import Document

from reports.models import Report, ReportImage, ReportPerson, Setup
from reports.services.report_render import date_range, render_report
from reports.tests.test_create_report import management, post_data
from reports.tests.test_phase2 import MediaMixin, png


def cover_cell_lines(doc, row, col):
    return [p.text for p in doc.tables[0].rows[row].cells[col].paragraphs if p.text.strip()]


class DateRangeTests(SimpleTestCase):
    def test_formats(self):
        d1, d2 = datetime.date(2026, 8, 13), datetime.date(2026, 8, 25)
        self.assertEqual(date_range(d1, d2), '8/13/2026 – 8/25/2026')
        self.assertEqual(date_range(d1, None), '8/13/2026')
        self.assertEqual(date_range(d1, d1), '8/13/2026')
        self.assertEqual(date_range(None, None), '')


class PersonnelEditorTests(TestCase):
    url = reverse('create-report')

    def test_people_saved_with_roles_in_order(self):
        data = post_data()
        data.update(management('people', 2))
        data.update({
            'people-0-name': 'Lucas Jennings', 'people-0-certification': 'Ultrasonic Level II',
            'people-0-prepared': 'on', 'people-0-examined': 'on',
            'people-1-name': 'Sky Tervo', 'people-1-reviewed': 'on',
        })
        self.client.post(self.url, data)
        people = list(Report.objects.get().people.values_list('name', 'prepared', 'examined', 'reviewed', 'order'))
        self.assertEqual(people, [('Lucas Jennings', True, True, False, 0), ('Sky Tervo', False, False, True, 1)])

    def test_person_needs_a_name(self):
        data = post_data()
        data.update(management('people', 1))
        data.update({'people-0-certification': 'Rope Access II', 'people-0-examined': 'on'})
        resp = self.client.post(self.url, data)
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(Report.objects.exists())

    def test_known_people_suggested_with_latest_certification(self):
        old = Report.objects.create()
        ReportPerson.objects.create(report=old, name='Monet White', certification='Rope access I')
        ReportPerson.objects.create(report=old, name='Monet White', certification='Rope access II')
        resp = self.client.get(self.url)
        self.assertIn(('Monet White', 'Rope access II'), resp.context['known_people'])
        self.assertContains(resp, 'data-certification="Rope access II"')

    def test_report_list_shows_prepared_by(self):
        report = Report.objects.create()
        ReportPerson.objects.create(report=report, name='A Prep', prepared=True, order=0)
        ReportPerson.objects.create(report=report, name='B Exam', examined=True, order=1)
        resp = self.client.get(reverse('report-list'))
        self.assertContains(resp, '<td>A Prep</td>')


class Phase3RenderTests(MediaMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.report = Report.objects.create(
            document_filename='r', test_date=datetime.date(2026, 8, 13), test_end_date=datetime.date(2026, 8, 25),
            comparison_title='HYDROFORM DATA COMPARISON (2022 TO 2026)',
        )
        Setup.objects.create(report=self.report, title='HydroFORM')
        for i, (name, cert, roles) in enumerate([
            ('Lucas Jennings', 'Ultrasonic Level II', ('prepared', 'examined')),
            ('Shawn Arangolord', 'Rope Access II', ('examined',)),
            ('Sky Tervo', '', ('reviewed',)),
        ]):
            ReportPerson.objects.create(report=self.report, order=i, name=name, certification=cert,
                                        **{role: True for role in roles})

    def render(self):
        return Document(io.BytesIO(render_report(self.report)))

    def test_cover_personnel_by_role(self):
        doc = self.render()
        self.assertEqual(cover_cell_lines(doc, 8, 3)[1:], ['Lucas Jennings\nUltrasonic Level II'])
        self.assertEqual(cover_cell_lines(doc, 9, 3)[1:], ['Lucas Jennings\nUltrasonic Level II', 'Shawn Arangolord\nRope Access II'])
        self.assertEqual(cover_cell_lines(doc, 10, 3)[1:], ['Sky Tervo'])

    def test_test_date_range_on_cover(self):
        self.assertEqual(cover_cell_lines(self.render(), 4, 5), ['8/13/2026 – 8/25/2026'])

    def test_comparison_grouped_under_titles(self):
        for i, (title, color) in enumerate([('Below CW4', 'red'), ('Below CW4', 'blue'), ('Below CW6', 'green')]):
            ReportImage.objects.create(report=self.report, kind=ReportImage.COMPARISON, caption=title,
                                       image=png(f'c{i}.png', color), order=i)
        doc = self.render()
        texts = [p.text for p in doc.paragraphs]
        self.assertIn('HYDROFORM DATA COMPARISON (2022 TO 2026)', texts)
        start = texts.index('HYDROFORM DATA COMPARISON (2022 TO 2026)')
        self.assertEqual([t for t in texts[start:start + 12] if t.startswith('Below')], ['Below CW4', 'Below CW6'])
        comparison_images = [s for s in doc.inline_shapes if round(s.width.inches, 1) == 7.0]
        self.assertEqual(len(comparison_images), 3)

    def test_comparison_section_left_out_when_empty(self):
        texts = [p.text for p in self.render().paragraphs]
        self.assertNotIn('HYDROFORM DATA COMPARISON (2022 TO 2026)', texts)
