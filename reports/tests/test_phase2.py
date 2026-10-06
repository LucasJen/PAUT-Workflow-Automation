"""Technique sections (title, procedure, calibration screenshots) and fixed results columns."""
import io
import shutil
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from docx import Document
from PIL import Image

from reports.models import Report, ResultsRow, ResultsTable, Setup, SetupImage
from reports.report_types import HIC_RESULTS_COLUMNS
from reports.results import fit_to_columns
from reports.services.report_render import render_report
from reports.tests.test_create_report import management, post_data

HIC = [heading for _, heading in HIC_RESULTS_COLUMNS]


def png(name, color='red'):
    buf = io.BytesIO()
    Image.new('RGB', (12, 8), color).save(buf, 'PNG')
    return SimpleUploadedFile(name, buf.getvalue(), content_type='image/png')


class FitToColumnsTests(SimpleTestCase):
    def test_matches_by_heading_ignoring_case_and_punctuation(self):
        cols, rows = fit_to_columns(['results', 'SCAN ID', 'Min Thk (in)'], [['ok', 'CW1', '0.8']], HIC)
        self.assertEqual(cols, HIC)
        self.assertEqual(rows[0][0], 'CW1')
        self.assertEqual(rows[0][5], '0.8')
        self.assertEqual(rows[0][6], 'ok')

    def test_unmatched_columns_fill_by_position(self):
        _, rows = fit_to_columns(['A', 'B'], [['CW1', 'Circ']], HIC)
        self.assertEqual(rows[0][:2], ['CW1', 'Circ'])

    def test_leftover_columns_are_kept_in_results_text(self):
        many = [f'C{i}' for i in range(9)]
        _, rows = fit_to_columns(many, [[str(i) for i in range(9)]], HIC)
        self.assertEqual(rows[0][6], '6\nC7: 7\nC8: 8')

    def test_no_fixed_columns_leaves_table_alone(self):
        self.assertEqual(fit_to_columns(['A'], [['1']], []), (['A'], [['1']]))


class MediaMixin:
    def setUp(self):
        media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, media, True)
        override = override_settings(MEDIA_ROOT=media)
        override.enable()
        self.addCleanup(override.disable)


class SetupScreenshotTests(MediaMixin, TestCase):
    url = reverse('create-report')

    def test_screenshots_saved_per_setup_and_removable(self):
        data = post_data(setups=[{'title': 'HydroFORM'}, {'title': 'Angle Beam'}])
        data['setups-0-cal_images'] = [png('a.png'), png('b.png', 'blue')]
        data['setups-1-cal_images'] = [png('c.png', 'green')]
        self.client.post(self.url, data)
        report = Report.objects.get()
        hydro, angle = report.setups.order_by('order')
        self.assertEqual(hydro.images.count(), 2)
        self.assertEqual(angle.images.count(), 1)

        # Remove one screenshot on the next save
        keep, drop = hydro.images.all()
        data = post_data(report=report, setups=[
            {'id': hydro.pk, 'report': report.pk, 'title': 'HydroFORM'},
            {'id': angle.pk, 'report': report.pk, 'title': 'Angle Beam'},
        ])
        data['remove_setup_image'] = [str(drop.pk)]
        self.client.post(self.url, data)
        self.assertEqual(list(hydro.images.values_list('pk', flat=True)), [keep.pk])

    def test_non_image_file_is_rejected_with_message(self):
        data = post_data(setups=[{'title': 'HydroFORM'}])
        data['setups-0-cal_images'] = [SimpleUploadedFile('notes.txt', b'not an image')]
        resp = self.client.post(self.url, data)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'is not an image')
        self.assertFalse(Report.objects.exists())

    def test_cannot_remove_another_reports_screenshot(self):
        other = Setup.objects.create(report=Report.objects.create())
        image = SetupImage.objects.create(setup=other, image=png('x.png'))
        data = post_data(setups=[{'title': 'T'}])
        data['remove_setup_image'] = [str(image.pk)]
        self.client.post(self.url, data)
        self.assertTrue(SetupImage.objects.filter(pk=image.pk).exists())


class TechniqueRenderTests(MediaMixin, TestCase):
    def test_sections_titles_procedures_and_screenshots(self):
        report = Report.objects.create(document_filename='r', procedure='FALLBACK-1')
        hydro = Setup.objects.create(report=report, order=0, title='HydroFORM', procedure='100-UT-031 Rev. 1')
        SetupImage.objects.create(setup=hydro, image=png('a.png'), order=0)
        SetupImage.objects.create(setup=hydro, image=png('b.png', 'blue'), order=1)
        Setup.objects.create(report=report, order=1, title='Angle Beam', procedure='100-UT-021 Rev 9.0')
        Setup.objects.create(report=report, order=2, title='TFM', procedure='100-UT-031 Rev. 1')  # duplicate procedure
        table = ResultsTable.objects.create(report=report, columns=['Results', 'Scan ID'])  # old, out-of-order table
        ResultsRow.objects.create(table=table, cells=['Blisters', 'CW1 Top'], order=0)

        doc = Document(io.BytesIO(render_report(report)))
        headings = [p.text for p in doc.paragraphs if p.text.startswith('Equipment Details:')]
        self.assertEqual(headings, ['Equipment Details: HydroFORM', 'Equipment Details: Angle Beam', 'Equipment Details: TFM'])

        text = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        cover = doc.tables[0]
        procedures = [p.text for p in cover.rows[10].cells[1].paragraphs if p.text]
        self.assertEqual(procedures, ['100-UT-031 Rev. 1', '100-UT-021 Rev 9.0'])  # once each, in order
        self.assertNotIn('FALLBACK-1', text)

        # Two screenshots under the HydroFORM table, sized two per line
        widths = [shape.width.inches for shape in doc.inline_shapes if round(shape.width.inches, 2) == 3.45]
        self.assertEqual(len(widths), 2)

        # Old table re-mapped to the fixed HIC columns by heading
        results = next(t for t in doc.tables if t.rows[0].cells[0].text.startswith('PAUT HydroFORM Work Scope'))
        self.assertEqual(results.rows[2].cells[0].text, 'CW1 Top')
        self.assertEqual(results.rows[2].cells[6].text, 'Blisters')

    def test_procedure_falls_back_to_report(self):
        report = Report.objects.create(procedure='100-UT-031 Rev. 1')
        Setup.objects.create(report=report, title='HydroFORM')
        doc = Document(io.BytesIO(render_report(report)))
        cover = doc.tables[0]
        self.assertEqual([p.text for p in cover.rows[10].cells[1].paragraphs if p.text], ['100-UT-031 Rev. 1'])


class EditorFixedColumnsTests(TestCase):
    def test_editor_gets_fixed_columns_and_fitted_rows(self):
        report = Report.objects.create()
        table = ResultsTable.objects.create(report=report, columns=['Results', 'Scan ID'])
        ResultsRow.objects.create(table=table, cells=['Blisters', 'CW1 Top'], order=0)
        resp = self.client.get(f"{reverse('create-report')}?loaded={report.pk}")
        self.assertEqual(resp.context['results_data']['columns'], HIC)
        self.assertEqual(resp.context['results_data']['rows'][0][0], 'CW1 Top')
        self.assertContains(resp, '"results_columns": [{"key": "scan_id", "heading": "Scan ID"}')
