import io
import os
import tempfile
from unittest import mock

from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from docx import Document

from reports.models import Report, Setup
from reports.views.reports import DOCX_CONTENT_TYPE, safe_filename


class SafeFilenameTests(SimpleTestCase):
    def test_strips_directories(self):
        for name in ['../../evil', '..\\..\\evil', 'C:\\Windows\\evil', '/etc/evil']:
            result = safe_filename(name)
            self.assertNotIn('/', result)
            self.assertNotIn('\\', result)
            self.assertNotIn(':', result)
            self.assertFalse(result.startswith('.'))

    def test_keeps_normal_names(self):
        self.assertEqual(safe_filename('Unit 3 - Weld 12'), 'Unit 3 - Weld 12')

    def test_blank_and_reserved_names_fall_back(self):
        for name in ['', '   ', '...', 'CON', 'nul']:
            self.assertEqual(safe_filename(name), 'report')


class GenerateReportOutputTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        override = override_settings(REPORT_OUTPUT_DIR=self.tmp.name)
        override.enable()
        self.addCleanup(override.disable)

    def generate(self, report):
        return self.client.get(reverse('generate-report', args=[report.pk]))

    def test_output_stays_in_output_dir(self):
        report = Report.objects.create(document_filename='../escape')
        Setup.objects.create(report=report)
        self.generate(report)
        self.assertEqual(os.listdir(self.tmp.name), ['_escape.docx'])

    def test_download_is_a_filled_word_document(self):
        report = Report.objects.create(document_filename='Unit 3 - Weld 12', client='ACME Refining')
        Setup.objects.create(report=report, scope_model='OmniScan X3')
        resp = self.generate(report)

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], DOCX_CONTENT_TYPE)
        self.assertIn('attachment', resp['Content-Disposition'])
        self.assertIn('Unit 3 - Weld 12.docx', resp['Content-Disposition'])

        doc = Document(io.BytesIO(b''.join(resp.streaming_content)))
        text = '\n'.join(
            [p.text for p in doc.paragraphs]
            + [c.text for t in doc.tables for r in t.rows for c in r.cells]
        )
        self.assertIn('ACME Refining', text)
        self.assertIn('OmniScan X3', text)
        self.assertNotIn('{{CLIENT}}', text)
        self.assertNotIn('{{SCOPE_MODEL}}', text)

    def test_server_copy_matches_download(self):
        report = Report.objects.create(document_filename='copy')
        Setup.objects.create(report=report)
        downloaded = b''.join(self.generate(report).streaming_content)
        with open(os.path.join(self.tmp.name, 'copy.docx'), 'rb') as f:
            self.assertEqual(f.read(), downloaded)

    def test_server_copy_can_be_turned_off(self):
        report = Report.objects.create(document_filename='nocopy')
        Setup.objects.create(report=report)
        with override_settings(REPORT_OUTPUT_DIR=None):
            resp = self.generate(report)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(os.listdir(self.tmp.name), [])

    def test_locked_server_copy_still_downloads(self):
        """If the previous copy is open in Word, the download must still work."""
        report = Report.objects.create(document_filename='locked')
        Setup.objects.create(report=report)
        real_open = open

        def locked_open(path, *args, **kwargs):
            if str(path).endswith('locked.docx'):
                raise PermissionError('file in use')
            return real_open(path, *args, **kwargs)

        with mock.patch('builtins.open', side_effect=locked_open), self.assertLogs('reports.views.reports', 'WARNING'):
            resp = self.generate(report)
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment', resp['Content-Disposition'])


class SaveAndDownloadTests(TestCase):
    url = reverse('create-report')

    def post(self, report, **extra):
        data = {
            'report_id': report.pk, 'setups-TOTAL_FORMS': '0', 'setups-INITIAL_FORMS': '0',
            'images-TOTAL_FORMS': '0', 'images-INITIAL_FORMS': '0',
            'drawings-TOTAL_FORMS': '0', 'drawings-INITIAL_FORMS': '0',
            'people-TOTAL_FORMS': '0', 'people-INITIAL_FORMS': '0',
            'comparison-TOTAL_FORMS': '0', 'comparison-INITIAL_FORMS': '0',
        }
        data.update(extra)
        return self.client.post(self.url, data)

    def test_save_and_download_returns_to_editor_which_starts_download(self):
        report = Report.objects.create()
        Setup.objects.create(report=report)
        resp = self.post(report, generate='')
        self.assertRedirects(resp, f'{self.url}?loaded={report.pk}&download=1')

        page = self.client.get(resp['Location'])
        self.assertContains(page, 'data-auto-download')
        self.assertContains(page, f'href="{reverse("generate-report", args=[report.pk])}"')

    def test_save_and_download_without_setups_explains(self):
        report = Report.objects.create()
        resp = self.client.get(self.post(report, generate='')['Location'])
        self.assertContains(resp, 'Add at least one UT setup')
        self.assertNotContains(resp, 'id="download-link"')

    def test_report_list_download_button_only_with_setups(self):
        with_setup = Report.objects.create()
        Setup.objects.create(report=with_setup)
        without = Report.objects.create()
        resp = self.client.get(reverse('report-list'))
        self.assertContains(resp, reverse('generate-report', args=[with_setup.pk]))
        self.assertNotContains(resp, reverse('generate-report', args=[without.pk]))
